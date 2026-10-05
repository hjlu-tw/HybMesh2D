#!/usr/bin/env python3
"""Build the JSST2026 conference deck (20 min) in two formats from ONE content source.

Outputs
    JSST2026_slides.pptx   editable PowerPoint, 16:9
    index.html             self-contained HTML deck (images embedded, arrow-key nav)

Content is transcribed from JSST2026_v1.pdf / .docx (2026-07-06). Every figure quoted
on a slide is traceable to a sentence in that paper; nothing is rounded up.
"""
import base64
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
PPTX_OUT = os.path.join(HERE, "JSST2026_slides.pptx")
HTML_OUT = os.path.join(HERE, "index.html")

# ---------------------------------------------------------------- palette
INK = RGBColor(0x14, 0x22, 0x2E)
NAVY = RGBColor(0x0B, 0x35, 0x52)
TEAL = RGBColor(0x10, 0x7A, 0x84)
AMBER = RGBColor(0xC2, 0x6A, 0x0B)
GREY = RGBColor(0x5B, 0x6B, 0x78)
MIST = RGBColor(0xF2, 0xF6, 0xF8)
LINE = RGBColor(0xD5, 0xDF, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX = {
    "ink": "#14222e", "navy": "#0b3552", "teal": "#107a84",
    "amber": "#c26a0b", "grey": "#5b6b78", "mist": "#f2f6f8",
    "line": "#d5dfe4", "white": "#ffffff",
}

FONT = "Arial"
SW, SH = 13.333, 7.5          # slide size, inches
M = 0.72                       # side margin
CW = SW - 2 * M                # content width

# ---------------------------------------------------------------- content
# Each slide: kind + its fields. notes = speaker script.
SLIDES = [
    dict(
        kind="title",
        eyebrow="JSST 2026",
        title="An Integrated CFD Workflow Platform\nfor In-House Solvers",
        subtitle="Modular Architecture and AI-Assisted Implementation",
        authors="Hsueh-Jui Lu,  Sheng-Jer Chen,  Fang-An Kuo",
        affil="National Center for High-performance Computing\n"
              "National Institutes of Applied Research, Taiwan",
        notes=(
            "Good morning. I'm Hsueh-Jui Lu, from the National Center for High-performance "
            "Computing in Taiwan.\n\n"
            "This talk is about a problem that most groups running an in-house CFD solver know "
            "well: the solver is good, and almost nobody outside the group can use it. We built "
            "a platform that closes that gap, and we built it with an AI assistant in the loop — "
            "so the talk is really two stories, the platform and the method.\n\n"
            "20 minutes. I'll leave time for questions."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="Motivation",
        heading="In-house solvers are numerically mature — and rarely adopted",
        lead="The barrier is the workflow, not the numerics.",
        bullets=[
            "One simulation means coordinating <b>five separate tools</b>",
            "Manual file conversions and inter-tool data transfers at every stage",
            "Solver-specific configuration, with little or no GUI support",
            "Even for a well-validated solver, familiarization takes <b>weeks</b>",
        ],
        notes=(
            "Start with the observation that motivates everything.\n\n"
            "In-house CFD solvers often reach high numerical maturity — they're validated, "
            "they're published, they produce good answers. And yet adoption stays low.\n\n"
            "The reason is almost never the numerics. It's that using one requires coordinating "
            "five separate tools, converting files by hand between every stage, and writing "
            "solver-specific configuration with essentially no GUI support.\n\n"
            "So the barrier is a workflow barrier. That's the thing we set out to remove."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="The representative case",
        heading="UNICONES",
        lead="A good solver with a familiarization cost measured in weeks.",
        bullets=[
            "Compressible Euler / Navier–Stokes solver based on the <b>CESE</b> method",
            "Developed at NCHC — National Institutes of Applied Research, Taiwan",
            "Openly available academically, yet still requires weeks of guided effort",
            "High numerical maturity, limited adoption — <i>exactly</i> the gap this work targets",
        ],
        notes=(
            "UNICONES is our in-house solver, and it's the representative case for the paper.\n\n"
            "It's a compressible Euler and Navier-Stokes solver built on the CESE method — the "
            "space-time conservation element and solution element framework.\n\n"
            "It's openly available academically. And even so, becoming productive with it takes "
            "on the order of weeks of guided effort. That combination — a good solver, a high "
            "adoption cost — is the gap we're addressing."
        ),
    ),
    dict(
        kind="cards",
        eyebrow="Contributions",
        heading="This paper contributes three things",
        cards=[
            dict(num="i", title="A modular solver-adapter interface",
                 body="Lightweight integration of heterogeneous solvers — lowering the adoption "
                      "barrier for new users."),
            dict(num="ii", title="A human–AI development methodology",
                 body="Documented, with empirical observations on how the contribution was "
                      "actually distributed."),
            dict(num="iii", title="A generalizable design strategy",
                 body="Workflow integration for research groups maintaining in-house solvers "
                      "with limited software-engineering resources."),
        ],
        notes=(
            "Three contributions, and I'll come back to each.\n\n"
            "First, the solver adapter — a small interface that lets you integrate a new solver "
            "without touching anything upstream.\n\n"
            "Second, the development methodology. We built this with an LLM assistant in agentic "
            "mode, and we kept enough records to say something concrete about what that was "
            "actually like — including where it helped least.\n\n"
            "Third, the design strategy itself, which we think generalizes to other groups in "
            "the same position."
        ),
    ),
    dict(
        kind="diagram",
        eyebrow="System architecture",
        heading="One environment, the whole pipeline",
        lead="Four modules, a single PyQt6 GUI, and automatic hand-off between every stage.",
        caption="Fig. 1  Data flow of the integrated CFD workflow platform.",
        notes=(
            "Here's the architecture. Four modules — PreProcessor, Mesh Generator, Core Solver, "
            "Results Visualization — inside one PyQt6 application.\n\n"
            "The important detail is the bar underneath. The Solver Adapter spans the last three "
            "modules, and it's the only place solver-specific knowledge lives. That's what makes "
            "the rest of the platform reusable.\n\n"
            "Every conversion between stages is automatic. The user never handles an intermediate "
            "file."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="Module 1",
        heading="PreProcessor — geometry preparation",
        bullets=[
            "Interactive geometry canvas",
            "Multiple point-spacing strategies",
            "Automatic sharp-corner detection",
            "Per-segment resampling",
            "Standalone C++ backend",
        ],
        notes=(
            "Module one, the preprocessor.\n\n"
            "It's an interactive canvas for building and editing geometry. You can choose among "
            "several point-spacing strategies, the tool detects sharp corners automatically, and "
            "you can resample individual segments rather than the whole curve.\n\n"
            "The heavy lifting sits in a standalone C++ backend, so the GUI stays responsive."
        ),
    ),
    dict(
        kind="figure",
        eyebrow="Module 2",
        heading="Mesh Generator — 2D hybrid meshes",
        figure="ogrid.png",
        caption="Structured boundary-layer quads growing from the body, unstructured "
                "triangles filling the far field.",
        bullets=[
            "Boundary-layer quads by <b>node advancement</b>",
            "Far-field triangles via Gmsh <b>Frontal-Delaunay</b>",
            "~1,500 lines of C++",
            "Exports VTK and STAR-CD (.vrt/.cel/.bnd)",
            "Inter-stage conversions handled automatically",
        ],
        notes=(
            "The mesh generator produces 2D hybrid meshes — about 1,500 lines of C++.\n\n"
            "You can see the structure in the figure: a structured boundary layer of "
            "quadrilaterals growing outward from the body by node advancement, and unstructured "
            "triangles filling the far field using Gmsh's Frontal-Delaunay algorithm.\n\n"
            "It exports both VTK and STAR-CD formats, and — this matters for the workflow claim — "
            "the conversions between stages are automatic."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="Module 3",
        heading="Core Solver — execution and monitoring",
        bullets=[
            "Input format conversion",
            "Optional MPI domain decomposition",
            "Executes solvers distributed as shared libraries, via a built-in DLL call template",
            "A background worker parses stdout in real time for <b>live convergence monitoring</b>",
        ],
        notes=(
            "Module three drives the solver.\n\n"
            "It handles input format conversion, optional MPI domain decomposition, and the "
            "execution itself — including solvers distributed as shared libraries, through a "
            "built-in DLL call template.\n\n"
            "The part users notice is the background worker: it parses stdout as the solve runs, "
            "so convergence is monitored live rather than after the fact."
        ),
    ),
    dict(
        kind="figure",
        eyebrow="Module 4",
        heading="Results Visualization",
        figure="naca_contour.png",
        caption="Cell-centred Mach number on a NACA 0012, rendered on the embedded canvas.",
        bullets=[
            "Parses Tecplot-format output",
            "Cell-centred scalar fields on an embedded matplotlib canvas",
            "Contours, streamlines and vector glyphs",
            "Interactive colormap and field-variable controls",
        ],
        notes=(
            "And module four, visualization.\n\n"
            "It parses Tecplot-format output and renders cell-centred scalar fields on an "
            "embedded matplotlib canvas — contours, streamlines, vector glyphs, with interactive "
            "control over the colormap and the field variable.\n\n"
            "This is a Mach number contour on a NACA 0012, produced end to end by the pipeline.\n\n"
            "The key design point: because it reads through the adapter, visualization is fully "
            "decoupled from any particular solver's output format."
        ),
    ),
    dict(
        kind="statement",
        eyebrow="The key idea",
        heading="The Solver Adapter",
        statement="Each adapter implements <b>three interface methods</b>. "
                  "Everything upstream stays unchanged.",
        stats=[
            ("3", "interface methods"),
            ("<150", "lines of Python per new solver"),
            ("0", "changes to upstream modules"),
        ],
        notes=(
            "Now the central design idea, and the one I'd like you to take away.\n\n"
            "Integrating a new solver means implementing three methods. That's the whole "
            "contract. Everything upstream — the preprocessor, the mesher, the visualization — "
            "stays exactly as it is.\n\n"
            "It comes to fewer than 150 lines of Python per new solver."
        ),
    ),
    dict(
        kind="table",
        eyebrow="The contract",
        heading="Three methods, one interface",
        table=dict(
            header=["Method", "Input", "Output", "Purpose"],
            rows=[
                ["convert_mesh(src, dst)", "VTK mesh path", "Solver-format mesh file(s)",
                 "Converts the platform mesh to the solver's input format — e.g. STAR-CD "
                 ".vrt/.cel/.bnd for UNICONES."],
                ["parse_stdout(line) → dict", "One line of solver stdout",
                 "Convergence quantities, or None",
                 "Extracts iteration, CFL and residual for live monitoring; returns None for "
                 "non-data lines."],
                ["read_results(path) → dataset", "Solver result file path",
                 "Structured field dataset",
                 "Parses solver-specific output for the visualization module, fully decoupling "
                 "it from the output format."],
            ],
        ),
        caption="Table 1  Solver Adapter interface definition.",
        notes=(
            "Here is the contract in full.\n\n"
            "convert_mesh takes the platform's VTK mesh and writes whatever the solver reads. "
            "For UNICONES that's the STAR-CD triplet.\n\n"
            "parse_stdout takes one line of solver stdout and returns convergence quantities, or "
            "None if it's not a data line. That's what drives live monitoring.\n\n"
            "read_results parses the solver's own output into a structured field dataset. Because "
            "visualization goes through this method, it never needs to know which solver ran.\n\n"
            "Three methods, and the solver is integrated."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="Methodology",
        heading="How it was built",
        bullets=[
            "LLM coding assistants in <b>agentic mode</b> — reading, writing and executing project "
            "files, not merely generating text",
            "Experts defined the pipeline architecture, the numerical algorithms, the "
            "format-compatibility chain, and physical validation",
            "AI generated GUI boilerplate, file parsers, and C++ routines from algorithmic "
            "descriptions",
            "Each feature converged in <b>2–5 rounds</b> of describe → generate → execute → "
            "validate → correct",
        ],
        notes=(
            "Now the second half of the talk: how this was built.\n\n"
            "We used LLM coding assistants in agentic mode — that means they could read, write "
            "and execute files in the project, not just produce text in a chat window. That "
            "distinction matters a great deal.\n\n"
            "The division of labour: the domain expert defined the architecture, the numerical "
            "algorithms, the chain of format compatibility, and what counted as physically "
            "correct. The AI generated GUI boilerplate, parsers, and C++ routines from "
            "algorithmic descriptions.\n\n"
            "Typically a feature took two to five rounds of describe, generate, execute, "
            "validate, correct."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="Methodology — a case in point",
        heading="Text was not enough",
        bullets=[
            "The expert described outward node advancement in natural language; the AI produced a "
            "first C++ implementation",
            "Text description alone <b>could not resolve the geometric failures</b>",
            "Actual runtime output was needed to find self-intersections at sharp trailing edges "
            "and wrong fan-node counts at convex corners",
            "More precise prompts consistently reduced the number of correction rounds",
            "Even a fix the AI reported as correct still needed <b>physical validation</b>, to "
            "catch discrepancies that appeared only in real use",
        ],
        notes=(
            "Let me make that concrete, because it's the most useful thing we learned.\n\n"
            "The boundary-layer module. The expert described the outward node-advancement "
            "algorithm in natural language, and the AI wrote a first C++ implementation. It "
            "compiled. It ran.\n\n"
            "And it was wrong in ways the description could not reach. Text alone could not "
            "resolve the geometric failures. What found them was actual runtime output — "
            "self-intersections at sharp trailing edges, incorrect fan-node counts at convex "
            "corners.\n\n"
            "Two things helped. More carefully phrased prompts consistently reduced the number of "
            "correction rounds. And even when the AI confirmed a fix as correct, physical "
            "validation by the domain expert was still needed — because some discrepancies only "
            "appear in real use."
        ),
    ),
    dict(
        kind="stats",
        eyebrow="Methodology — the record",
        heading="What the logs say",
        lead="Reported as case-study observations, not controlled experimental results.",
        stats=[
            ("~70–80%", "of code AI-generated, by line count"),
            ("2–5", "median correction rounds per feature"),
            ("~4–6 weeks", "total effort, vs. 3–6 months conventional"),
            ("1", "domain expert — no dedicated GUI engineer"),
        ],
        notes=(
            "Some numbers, from the developer logs. I want to be careful here: these are "
            "case-study observations, not controlled experimental results. Treat them as "
            "indicative.\n\n"
            "By line count, roughly 70 to 80 percent of the code was AI-generated.\n\n"
            "The median feature took two to five correction rounds.\n\n"
            "Total effort was around four to six weeks — against an estimated three to six months "
            "for conventional solo development of equivalent scope.\n\n"
            "And it was one domain expert. There was no dedicated GUI software engineer on this "
            "project."
        ),
    ),
    dict(
        kind="distribution",
        eyebrow="Methodology — where it helped",
        heading="The contribution was unevenly distributed",
        bars=[
            ("GUI boilerplate, widget layouts, signal–slot wiring, file I/O", 85, "~80–90%"),
            ("Cross-component integration and domain-specific numerical logic", 50, "~40–60%"),
        ],
        footnote="The 69-file, 18,000-line PyQt6 GUI was highly AI-productive for boilerplate — "
                 "but workflow-level decisions (mode switching, cross-panel consistency) needed "
                 "explicit human specification. The IBM and non-IBM output modes were discovered, "
                 "only in live testing, to produce different stdout structures.",
        notes=(
            "The most honest finding is that the contribution was uneven.\n\n"
            "For GUI boilerplate — widget layouts, signal-slot wiring, stdout parsers — the AI "
            "was extremely productive, 80 to 90 percent.\n\n"
            "For domain-specific numerical logic and cross-component integration, it was much "
            "lower — 40 to 60 percent — and those are exactly the parts that decide whether the "
            "software is right.\n\n"
            "The 69-file, 18,000-line GUI reinforced the same pattern. Boilerplate: very "
            "productive. Workflow-level decisions — how modes switch, how panels stay consistent "
            "— needed explicit human specification. A good example: the IBM and non-IBM solver "
            "output modes turned out, only in live testing, to produce different stdout "
            "structures. No amount of prompt-writing would have surfaced that; running it did."
        ),
    ),
    dict(
        kind="stats",
        eyebrow="Evaluation",
        heading="Before and after",
        lead="Eliminating manual format conversion and inter-tool transfer at every stage.",
        stats=[
            ("5 → 1", "separate tools required"),
            ("~20 → ~5", "manual steps per simulation"),
            ("~2 weeks → ~2 days", "time for a new user to become familiar"),
        ],
        footnote="Indicative figures from informal observation, not controlled experimental "
                 "results.",
        notes=(
            "So what did it buy?\n\n"
            "Against using UNICONES without a unified interface: five separate tools become one. "
            "Around twenty manual steps per simulation become around five.\n\n"
            "And the time for a new user to become familiar with the full workflow drops from "
            "roughly two weeks to roughly two days.\n\n"
            "Again — these are estimates from informal observation, not a controlled study. We "
            "say so in the paper. The mechanism behind them is concrete, though: it's the "
            "elimination of manual format conversion and inter-tool transfer at each stage."
        ),
    ),
    dict(
        kind="cards",
        eyebrow="Discussion",
        heading="The oversight that proved indispensable",
        lead="Three layers, and the method fails without all three.",
        cards=[
            dict(num="1", title="Precise problem description",
                 body="More careful phrasing consistently reduced correction rounds."),
            dict(num="2", title="Execution-based diagnosis",
                 body="Runtime output found the edge cases that description could not."),
            dict(num="3", title="Physical validation",
                 body="An expert still had to catch discrepancies visible only in real use."),
        ],
        footnote="Without domain expertise at each layer, AI-assisted scientific software "
                 "development risks producing <b>plausible but incorrect</b> implementations.",
        notes=(
            "Stepping back — the AI-assisted approach proved most effective for peripheral "
            "tooling: GUIs, data pipelines, visualization. That's where LLM strengths in "
            "well-established software patterns apply.\n\n"
            "But effective collaboration required three layers of human oversight: precise "
            "problem description, execution-based diagnosis of edge cases, and physical "
            "validation.\n\n"
            "And I'd underline the conclusion we draw: without domain expertise at each of those "
            "layers, this approach risks producing implementations that are plausible and "
            "incorrect. In CFD that's the worst possible failure mode, because it doesn't look "
            "like one."
        ),
    ),
    dict(
        kind="bullets",
        eyebrow="Conclusion",
        heading="What we take away",
        bullets=[
            "A unified CFD workflow platform for in-house solvers, demonstrated with UNICONES — "
            "without modifying the solver core",
            "Tools reduced <b>5 → 1</b>, workflow steps <b>~20 → ~5</b>, onboarding "
            "<b>~2 weeks → ~2 days</b>",
            "The Solver Adapter — <b>under 150 lines of Python per new solver</b> — decouples "
            "solver-specific logic from the platform core",
            "A generalizable strategy for groups maintaining in-house solvers with limited "
            "software-engineering resources",
        ],
        notes=(
            "To conclude.\n\n"
            "We presented an integrated CFD workflow platform for in-house solvers, demonstrated "
            "with UNICONES, built without modifying the solver core.\n\n"
            "It reduces the tools required from five to one, the workflow from around twenty "
            "steps to around five, and estimated onboarding from two weeks to two days.\n\n"
            "The modular adapter — under 150 lines of Python per new solver — decouples "
            "solver-specific logic from the platform core. And we think that pattern is "
            "reproducible by domain-expert teams without dedicated software-engineering "
            "resources."
        ),
    ),
    dict(
        kind="closing",
        eyebrow="Future work",
        heading="Where this goes next",
        bullets=[
            "Extend the platform to <b>3D mesh generation</b>",
            "A formal user study with standardized usability metrics",
            "Web-based deployment, to further lower the barrier for CFD education and research",
        ],
        thanks="Thank you — questions welcome.",
        affil="This work was made possible by the academic resources and advanced research "
              "infrastructure of the National Center for High-performance Computing, National "
              "Institutes of Applied Research, Taiwan, with partial funding from the fundamental "
              "R&D program of NSTC, Taiwan.",
        notes=(
            "Future work, briefly: extending to 3D mesh generation, a formal user study with "
            "standardized usability metrics — which is what our current estimates really need — "
            "and web-based deployment, to lower the barrier further for CFD education and "
            "research.\n\n"
            "Thank you. I'm happy to take questions."
        ),
    ),
]

# ---------------------------------------------------------------- pptx helpers


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def tb(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def run(p, text, size, color=INK, bold=False, italic=False, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return r


def rich(p, markup, size, color=INK, bold=False, font=FONT):
    """Render <b>/<i> markup into runs on paragraph p."""
    import re
    for chunk in re.split(r"(<b>.*?</b>|<i>.*?</i>)", markup):
        if not chunk:
            continue
        if chunk.startswith("<b>"):
            run(p, chunk[3:-4], size, color, bold=True, font=font)
        elif chunk.startswith("<i>"):
            run(p, chunk[3:-4], size, color, bold=bold, italic=True, font=font)
        else:
            run(p, chunk, size, color, bold=bold, font=font)


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        solid(s, fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    return s


def header(slide, eyebrow, heading, lead=None):
    """Standard slide header. Returns y where body content may start."""
    if eyebrow:
        tf = tb(slide, M, 0.52, CW, 0.28)
        run(tf.paragraphs[0], eyebrow.upper(), 11.5, TEAL, bold=True)
    tf = tb(slide, M, 0.86, CW, 0.7)
    run(tf.paragraphs[0], heading, 27, NAVY, bold=True)
    y = 1.62
    if lead:
        tf = tb(slide, M, y, CW * 0.86, 0.4)
        rich(tf.paragraphs[0], lead, 14, GREY)
        y += 0.46
    return y


def bullet_list(slide, x, y, w, bullets, size=15, gap=0.52, color=INK):
    tf = tb(slide, x, y, w, 4.5)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        run(p, "—  ", size, TEAL, bold=True)
        rich(p, b, size, color)
        if i < len(bullets) - 1:
            p.space_after = Pt(11)
    return tf


def footer(slide, n):
    tf = tb(slide, M, SH - 0.52, CW * 0.7, 0.24)
    run(tf.paragraphs[0], "An Integrated CFD Workflow Platform for In-House Solvers",
        8.5, RGBColor(0x9A, 0xA8, 0xB0))
    tf = tb(slide, SW - M - 0.8, SH - 0.52, 0.8, 0.24, align=PP_ALIGN.RIGHT)
    run(tf.paragraphs[0], str(n), 9.5, RGBColor(0x9A, 0xA8, 0xB0), bold=True)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- pptx renderers


def s_title(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SW, SH, fill=NAVY)
    rect(slide, 0, 0, 0.22, SH, fill=TEAL)

    tf = tb(slide, 1.15, 1.62, 9, 0.3)
    run(tf.paragraphs[0], s["eyebrow"].upper(), 13, RGBColor(0x7F, 0xC6, 0xCE), bold=True)

    tf = tb(slide, 1.15, 2.08, 11, 1.9)
    for i, line in enumerate(s["title"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run(p, line, 40, WHITE, bold=True)

    tf = tb(slide, 1.15, 3.62, 10.5, 0.4)
    run(tf.paragraphs[0], s["subtitle"], 18, RGBColor(0x9F, 0xC5, 0xD6))

    rect(slide, 1.15, 4.30, 1.5, 0.035, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)

    tf = tb(slide, 1.15, 4.62, 10, 0.34)
    run(tf.paragraphs[0], s["authors"], 15, WHITE, bold=True)

    tf = tb(slide, 1.15, 5.04, 10, 0.8)
    for i, line in enumerate(s["affil"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run(p, line, 12.5, RGBColor(0x8F, 0xA8, 0xB8))
    add_notes(slide, s["notes"])
    return slide


def s_bullets(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"], s.get("lead"))
    bullet_list(slide, M, y + 0.25, CW * 0.92, s["bullets"])
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_cards(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"], s.get("lead"))
    cards = s["cards"]
    gap = 0.34
    w = (CW - gap * (len(cards) - 1)) / len(cards)
    top = y + 0.16
    bot = SH - 0.86 - (0.78 if s.get("footnote") else 0.0)
    h = bot - top
    # vertically centre the text block inside each card
    inner_h = 0.42 + 0.82 + 0.96          # num + title + body allowance
    pad = max(0.30, (h - inner_h) / 2.0)
    for i, c in enumerate(cards):
        x = M + i * (w + gap)
        rect(slide, x, top, w, h, fill=MIST, line=LINE)
        rect(slide, x, top, w, 0.05, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
        tf = tb(slide, x + 0.3, top + pad, w - 0.6, 0.42)
        run(tf.paragraphs[0], c["num"], 22, TEAL, bold=True)
        tf = tb(slide, x + 0.3, top + pad + 0.48, w - 0.6, 0.82)
        run(tf.paragraphs[0], c["title"], 15.5, NAVY, bold=True)
        tf = tb(slide, x + 0.3, top + pad + 1.30, w - 0.6, 0.96)
        rich(tf.paragraphs[0], c["body"], 12, GREY)
    if s.get("footnote"):
        tf = tb(slide, M, bot + 0.30, CW, 0.6)
        rich(tf.paragraphs[0], s["footnote"], 12.5, INK)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_statement(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"])
    tf = tb(slide, M, y + 0.30, CW * 0.94, 1.0)
    rich(tf.paragraphs[0], s["statement"], 20, INK)
    stats = s["stats"]
    gap = 0.34
    w = (CW - gap * (len(stats) - 1)) / len(stats)
    ty = y + 1.85
    for i, (big, lab) in enumerate(stats):
        x = M + i * (w + gap)
        rect(slide, x, ty, w, 1.6, fill=MIST, line=LINE)
        tf = tb(slide, x + 0.28, ty + 0.30, w - 0.56, 0.7)
        run(tf.paragraphs[0], big, 34, TEAL, bold=True)
        tf = tb(slide, x + 0.28, ty + 1.06, w - 0.56, 0.5)
        run(tf.paragraphs[0], lab, 12, GREY)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_stats(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"], s.get("lead"))
    stats = s["stats"]
    if len(stats) == 4:
        cols, rows, ch = 2, 2, 1.28
    else:
        cols, rows, ch = 3, 1, 1.86
    gap = 0.34
    w = (CW - gap * (cols - 1)) / cols
    # numbers that wrap to two lines must not push their label down out of line,
    # so the number gets a fixed two-line slot and the label sits at a fixed y
    big_h = 0.96
    for i, (big, lab) in enumerate(stats):
        r, c = divmod(i, cols)
        x = M + c * (w + gap)
        yy = y + 0.32 + r * (ch + 0.30)
        rect(slide, x, yy, w, ch, fill=MIST, line=LINE)
        pad = (ch - big_h - 0.42) / 2.0
        tf = tb(slide, x + 0.32, yy + pad, w - 0.64, big_h)
        run(tf.paragraphs[0], big, 30, TEAL, bold=True)
        tf = tb(slide, x + 0.32, yy + pad + big_h + 0.06, w - 0.64, 0.42)
        run(tf.paragraphs[0], lab, 12.5, GREY)
    if s.get("footnote"):
        tf = tb(slide, M, y + 0.32 + rows * (ch + 0.30) + 0.12, CW, 0.6)
        rich(tf.paragraphs[0], s["footnote"], 12.5, INK)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_distribution(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"])
    yy = y + 0.42
    for label, pct, tag in s["bars"]:
        tf = tb(slide, M, yy, CW - 1.3, 0.34)
        run(tf.paragraphs[0], label, 13.5, INK)
        tf = tb(slide, SW - M - 1.15, yy - 0.04, 1.15, 0.4, align=PP_ALIGN.RIGHT)
        run(tf.paragraphs[0], tag, 15, AMBER, bold=True)
        bar_y = yy + 0.44
        rect(slide, M, bar_y, CW - 1.3, 0.34, fill=MIST, shape=MSO_SHAPE.RECTANGLE)
        rect(slide, M, bar_y, (CW - 1.3) * pct / 100.0, 0.34, fill=TEAL,
             shape=MSO_SHAPE.RECTANGLE)
        yy += 1.30
    tf = tb(slide, M, yy + 0.18, CW * 0.96, 1.6)
    rich(tf.paragraphs[0], s["footnote"], 12.5, GREY)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_figure(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"])
    from PIL import Image
    path = os.path.join(FIG, s["figure"])
    iw, ih = Image.open(path).size
    box_h = 3.85
    box_w = 6.5
    scale = min(box_w / iw, box_h / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(path, Inches(M), Inches(y + 0.24 + (box_h - h) / 2),
                             width=Inches(w), height=Inches(h))
    if s.get("caption"):
        tf = tb(slide, M, y + 0.24 + box_h + 0.06, box_w, 0.5)
        run(tf.paragraphs[0], s["caption"], 10.5, GREY, italic=True)
    bx = M + box_w + 0.5
    bullet_list(slide, bx, y + 0.45, SW - M - bx, s["bullets"], size=13.5, gap=0.5)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_table(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"])
    t = s["table"]
    rows, cols = len(t["rows"]) + 1, len(t["header"])
    width = Inches(CW)
    height = Inches(0.5 * rows)
    gt = slide.shapes.add_table(rows, cols, Inches(M), Inches(y + 0.22),
                                width, height).table
    for c, wfrac in enumerate([0.235, 0.175, 0.185, 0.405]):
        gt.columns[c].width = Emu(int(Inches(CW) * wfrac))
    gt.rows[0].height = Inches(0.42)
    for r in range(1, rows):
        gt.rows[r].height = Inches(0.86)
    for c, htxt in enumerate(t["header"]):
        cell = gt.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = cell.margin_right = Inches(0.09)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        run(p, htxt, 12, WHITE, bold=True)
    for r, row in enumerate(t["rows"], start=1):
        for c, val in enumerate(row):
            cell = gt.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else MIST
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            mono = (c == 0)
            run(p, val, 11.5 if c else 10.5,
                TEAL if c == 0 else INK,
                bold=mono, font="Consolas" if mono else FONT)
    if s.get("caption"):
        tf = tb(slide, M, y + 0.22 + 0.5 * rows + 0.14, CW, 0.4)
        run(tf.paragraphs[0], s["caption"], 10.5, GREY, italic=True)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_diagram(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, s["eyebrow"], s["heading"], s.get("lead"))

    mods = [("PreProcessor", "Geometry & surface points", "→ .dat / .csv"),
            ("Mesh Generator", "Hybrid mesh generation", "→ VTK + STAR-CD"),
            ("Core Solver", "Execution + monitoring", "→ Tecplot format"),
            ("Results Visualization", "Tecplot parser", "→ graph / contour")]
    top = y + 0.30
    bh = 1.62
    gap = 0.42
    bw = (CW - gap * 3) / 4
    for i, (name, sub, out) in enumerate(mods):
        x = M + i * (bw + gap)
        rect(slide, x, top, bw, bh, fill=MIST, line=LINE)
        rect(slide, x, top, bw, 0.055, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
        tf = tb(slide, x + 0.18, top + 0.30, bw - 0.36, 0.5, align=PP_ALIGN.CENTER)
        run(tf.paragraphs[0], name, 14.5, NAVY, bold=True)
        tf = tb(slide, x + 0.14, top + 0.78, bw - 0.28, 0.4, align=PP_ALIGN.CENTER)
        run(tf.paragraphs[0], sub, 10.5, GREY)
        tf = tb(slide, x + 0.14, top + 1.18, bw - 0.28, 0.34, align=PP_ALIGN.CENTER)
        run(tf.paragraphs[0], out, 10, TEAL)
        if i < 3:
            ar = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + bw + 0.07), Inches(top + bh / 2 - 0.10),
                Inches(gap - 0.14), Inches(0.20))
            solid(ar, RGBColor(0xB8, 0xC9, 0xD2))
            ar.shadow.inherit = False

    ay = top + bh + 0.52
    ax = M + (bw + gap)
    aw = CW - (bw + gap)
    rect(slide, ax, ay, aw, 1.72, fill=RGBColor(0xE8, 0xF2, 0xF3), line=RGBColor(0xA9, 0xD3, 0xD8))
    tf = tb(slide, ax + 0.30, ay + 0.22, aw - 0.6, 0.36)
    run(tf.paragraphs[0], "Solver Adapter", 14, TEAL, bold=True)
    pills = ["convert_mesh()", "parse_stdout()", "read_results()"]
    pw = (aw - 0.6 - 0.28 * 2) / 3
    for i, pl in enumerate(pills):
        px = ax + 0.30 + i * (pw + 0.28)
        rect(slide, px, ay + 0.76, pw, 0.62, fill=WHITE, line=RGBColor(0xA9, 0xD3, 0xD8))
        tf = tb(slide, px, ay + 0.76, pw, 0.62, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        run(tf.paragraphs[0], pl, 11.5, NAVY, bold=True, font="Consolas")

    if s.get("caption"):
        tf = tb(slide, M, SH - 0.86, CW, 0.3)
        run(tf.paragraphs[0], s["caption"], 10.5, GREY, italic=True)
    footer(slide, n)
    add_notes(slide, s["notes"])
    return slide


def s_closing(prs, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SW, SH, fill=NAVY)
    rect(slide, 0, 0, 0.22, SH, fill=TEAL)
    tf = tb(slide, 1.1, 1.05, 10, 0.3)
    run(tf.paragraphs[0], s["eyebrow"].upper(), 12.5, RGBColor(0x7F, 0xC6, 0xCE), bold=True)
    tf = tb(slide, 1.1, 1.45, 11, 0.7)
    run(tf.paragraphs[0], s["heading"], 30, WHITE, bold=True)
    tf = tb(slide, 1.1, 2.55, 10.6, 1.9)
    for i, b in enumerate(s["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        run(p, "—  ", 15, RGBColor(0x7F, 0xC6, 0xCE), bold=True)
        rich(p, b, 15, RGBColor(0xE4, 0xEE, 0xF3))
    rect(slide, 1.1, 4.72, 1.5, 0.035, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    tf = tb(slide, 1.1, 5.02, 11, 0.4)
    run(tf.paragraphs[0], s["thanks"], 18, WHITE, bold=True)
    tf = tb(slide, 1.1, 5.66, 10.8, 1.1)
    run(tf.paragraphs[0], s["affil"], 10.5, RGBColor(0x8F, 0xA8, 0xB8))
    add_notes(slide, s["notes"])
    return slide


RENDER = {
    "title": s_title, "bullets": s_bullets, "cards": s_cards,
    "statement": s_statement, "stats": s_stats, "distribution": s_distribution,
    "figure": s_figure, "table": s_table, "diagram": s_diagram, "closing": s_closing,
}


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    for i, s in enumerate(SLIDES, start=1):
        RENDER[s["kind"]](prs, s, i)
    prs.save(PPTX_OUT)
    print(f"pptx  {PPTX_OUT}  ({len(SLIDES)} slides)")


# ---------------------------------------------------------------- html


def b64(name):
    with open(os.path.join(FIG, name), "rb") as f:
        return "data:image/png;base64," + base64.b64encode(f.read()).decode()


def h(markup):
    return markup.replace("<b>", "<strong>").replace("</b>", "</strong>") \
                 .replace("<i>", "<em>").replace("</i>", "</em>")


def html_body():
    out = []
    for i, s in enumerate(SLIDES, start=1):
        k = s["kind"]
        cls = f"slide slide--{k}"
        parts = [f'<section class="{cls}" data-n="{i}">']
        if k == "title":
            parts.append(f'''<div class="title-wrap">
  <p class="eyebrow eyebrow--dark">{h(s["eyebrow"])}</p>
  <h1>{h(s["title"]).replace(chr(10), "<br>")}</h1>
  <p class="subtitle">{h(s["subtitle"])}</p>
  <hr class="rule rule--light">
  <p class="authors">{h(s["authors"])}</p>
  <p class="affil">{h(s["affil"]).replace(chr(10), "<br>")}</p>
</div>''')
        elif k == "closing":
            parts.append(f'''<div class="closing-wrap">
  <p class="eyebrow">{h(s["eyebrow"])}</p>
  <h1 class="h1--light">{h(s["heading"])}</h1>
  <ul class="bullets bullets--light">
    {''.join(f"<li>{h(b)}</li>" for b in s["bullets"])}
  </ul>
  <hr class="rule rule--light">
  <p class="thanks">{h(s["thanks"])}</p>
  <p class="ack">{h(s["affil"])}</p>
</div>''')
        else:
            head = ['<header class="slide-head">']
            if s.get("eyebrow"):
                head.append(f'<p class="eyebrow">{h(s["eyebrow"])}</p>')
            head.append(f'<h2>{h(s["heading"])}</h2>')
            if s.get("lead"):
                head.append(f'<p class="lead">{h(s["lead"])}</p>')
            head.append("</header>")
            parts.append("".join(head))
            parts.append('<div class="slide-body">')

            if k == "bullets":
                parts.append('<ul class="bullets">' +
                             "".join(f"<li>{h(b)}</li>" for b in s["bullets"]) + "</ul>")
            elif k == "cards":
                parts.append('<div class="cards">' + "".join(
                    f'<article class="card"><span class="card-num">{c["num"]}</span>'
                    f'<h3>{h(c["title"])}</h3><p>{h(c["body"])}</p></article>'
                    for c in s["cards"]) + "</div>")
                if s.get("footnote"):
                    parts.append(f'<p class="footnote">{h(s["footnote"])}</p>')
            elif k == "statement":
                parts.append(f'<p class="statement">{h(s["statement"])}</p>')
                parts.append('<div class="stats stats--row">' + "".join(
                    f'<div class="stat"><span class="stat-big">{big}</span>'
                    f'<span class="stat-lab">{lab}</span></div>'
                    for big, lab in s["stats"]) + "</div>")
            elif k == "stats":
                parts.append('<div class="stats stats--%s">' % (
                    "grid" if len(s["stats"]) == 4 else "row") + "".join(
                    f'<div class="stat"><span class="stat-big">{big}</span>'
                    f'<span class="stat-lab">{lab}</span></div>'
                    for big, lab in s["stats"]) + "</div>")
                if s.get("footnote"):
                    parts.append(f'<p class="footnote">{h(s["footnote"])}</p>')
            elif k == "distribution":
                parts.append('<div class="bars">' + "".join(
                    f'<div class="bar-row"><div class="bar-top"><span class="bar-lab">{h(label)}</span>'
                    f'<span class="bar-tag">{tag}</span></div>'
                    f'<div class="bar-track"><div class="bar-fill" style="width:{pct}%"></div></div>'
                    f'</div>' for label, pct, tag in s["bars"]) + "</div>")
                parts.append(f'<p class="footnote">{h(s["footnote"])}</p>')
            elif k == "figure":
                parts.append(f'''<div class="figure-split">
  <figure><img src="{b64(s["figure"])}" alt="">
    <figcaption>{h(s.get("caption",""))}</figcaption></figure>
  <ul class="bullets bullets--tight">{''.join(f"<li>{h(b)}</li>" for b in s["bullets"])}</ul>
</div>''')
            elif k == "table":
                t = s["table"]
                parts.append('<table class="adapter"><thead><tr>' +
                             "".join(f"<th>{h(c)}</th>" for c in t["header"]) +
                             "</tr></thead><tbody>" +
                             "".join("<tr>" + "".join(
                                 f'<td class="{"mono" if j == 0 else ""}">{h(v)}</td>'
                                 for j, v in enumerate(row)) + "</tr>" for row in t["rows"]) +
                             "</tbody></table>")
                parts.append(f'<p class="footnote">{h(s["caption"])}</p>')
            elif k == "diagram":
                mods = [("PreProcessor", "Geometry &amp; surface points", "→ .dat / .csv"),
                        ("Mesh Generator", "Hybrid mesh generation", "→ VTK + STAR-CD"),
                        ("Core Solver", "Execution + monitoring", "→ Tecplot format"),
                        ("Results Visualization", "Tecplot parser", "→ graph / contour")]
                parts.append('<div class="flow">')
                for j, (nm, sub, o) in enumerate(mods):
                    parts.append(f'<div class="flow-mod"><h3>{nm}</h3><p>{sub}</p>'
                                 f'<span class="flow-out">{o}</span></div>')
                    if j < 3:
                        parts.append('<div class="flow-arrow" aria-hidden="true">→</div>')
                parts.append("</div>")
                parts.append('<div class="adapter-bar"><span class="adapter-name">Solver Adapter</span>'
                             '<div class="pills">' +
                             "".join(f'<span class="pill">{p}</span>' for p in
                                     ["convert_mesh()", "parse_stdout()", "read_results()"]) +
                             "</div></div>")
                parts.append(f'<p class="footnote">{h(s["caption"])}</p>')
            parts.append("</div>")

        parts.append('<div class="slide-num"></div>')
        parts.append("</section>")
        out.append("".join(parts))
    return "\n".join(out)


HTML = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>An Integrated CFD Workflow Platform for In-House Solvers</title>
<style>
:root{
  --ink:#14222e; --navy:#0b3552; --teal:#107a84; --amber:#c26a0b;
  --grey:#5b6b78; --mist:#f2f6f8; --line:#d5dfe4; --white:#fff;
  --font:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Helvetica Neue",Arial,sans-serif;
  --mono:ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;
}
*{box-sizing:border-box;margin:0;padding:0}
html,body{height:100%}
body{background:#0e1720;font-family:var(--font);color:var(--ink);
  -webkit-font-smoothing:antialiased;overflow:hidden}
#deck{position:fixed;inset:0;display:grid;place-items:center}
.slide{position:absolute;width:1280px;height:720px;background:var(--white);
  padding:52px 69px 46px;display:none;flex-direction:column;overflow:hidden;
  transform-origin:center center}
.slide.is-active{display:flex}
/* header */
.slide-head{margin-bottom:22px;flex:none}
.eyebrow{font-size:11.5px;font-weight:700;letter-spacing:.13em;text-transform:uppercase;
  color:var(--teal);margin-bottom:9px}
.eyebrow--dark{color:#7fc6ce}
h2{font-size:31px;font-weight:700;color:var(--navy);line-height:1.18;letter-spacing:-.012em}
.lead{font-size:15.5px;color:var(--grey);margin-top:11px;line-height:1.5;max-width:88%}
.slide-body{flex:1;display:flex;flex-direction:column;min-height:0;justify-content:center}
/* bullets */
.bullets{list-style:none;display:flex;flex-direction:column;gap:17px;max-width:94%}
.bullets li{position:relative;padding-left:26px;font-size:17.5px;line-height:1.5}
.bullets li::before{content:"—";position:absolute;left:0;color:var(--teal);font-weight:700}
.bullets--tight li{font-size:15px;gap:0}
.bullets--tight{max-width:100%;gap:14px}
.bullets--light li{color:#e4eef3}
.bullets--light li::before{color:#7fc6ce}
strong{font-weight:700}
/* cards */
.cards{display:grid;grid-template-columns:repeat(var(--n,3),1fr);gap:26px;flex:1}
.card{background:var(--mist);border:1px solid var(--line);border-top:3px solid var(--teal);
  padding:30px 26px;display:flex;flex-direction:column;gap:12px;justify-content:center}
.card-num{font-size:22px;font-weight:700;color:var(--teal);line-height:1}
.card h3{font-size:17px;color:var(--navy);line-height:1.3}
.card p{font-size:13.5px;color:var(--grey);line-height:1.55}
/* stats */
.stats{display:grid;gap:26px;margin-top:8px}
.stats--row{grid-template-columns:repeat(3,1fr)}
.stats--grid{grid-template-columns:repeat(2,1fr)}
.stat{background:var(--mist);border:1px solid var(--line);padding:26px 26px;
  display:flex;flex-direction:column;gap:8px}
.stat-big{font-size:38px;font-weight:700;color:var(--teal);line-height:1.06;letter-spacing:-.02em;
  display:block;min-height:2.16em}
.stat-lab{font-size:13.5px;color:var(--grey);line-height:1.45}
.statement{font-size:22px;line-height:1.5;margin-bottom:26px;max-width:94%}
/* bars */
.bars{display:flex;flex-direction:column;gap:40px}
.bar-top{display:flex;justify-content:space-between;align-items:baseline;margin-bottom:11px}
.bar-lab{font-size:15.5px}
.bar-tag{font-size:17px;font-weight:700;color:var(--amber)}
.bar-track{height:26px;background:var(--mist);border:1px solid var(--line)}
.bar-fill{height:100%;background:var(--teal)}
/* figure */
.figure-split{display:grid;grid-template-columns:1.28fr .72fr;
  grid-template-rows:minmax(0,1fr);gap:38px;flex:1;min-height:0;align-items:center}
.figure-split figure{display:flex;flex-direction:column;min-height:0;height:100%;
  justify-content:center}
.figure-split img{max-width:100%;max-height:100%;object-fit:contain;
  border:1px solid var(--line);align-self:flex-start}
.figure-split .bullets{max-width:100%}
figcaption{font-size:11px;color:var(--grey);font-style:italic;margin-top:10px;line-height:1.45}
/* table */
.adapter{width:100%;border-collapse:collapse;font-size:13.5px;margin-top:4px}
.adapter th{background:var(--navy);color:#fff;text-align:left;padding:12px 14px;
  font-size:12.5px;font-weight:700}
.adapter td{padding:16px 14px;border-bottom:1px solid var(--line);vertical-align:middle;
  line-height:1.5}
.adapter tr:nth-child(even) td{background:var(--mist)}
.adapter .mono{font-family:var(--mono);font-size:12.5px;font-weight:700;color:var(--teal);
  white-space:nowrap}
/* flow diagram */
.flow{display:flex;align-items:center;gap:0;margin-top:14px}
.flow-mod{flex:1;background:var(--mist);border:1px solid var(--line);
  border-top:3px solid var(--teal);padding:22px 14px;text-align:center;
  display:flex;flex-direction:column;gap:7px;min-height:158px;justify-content:center}
.flow-mod h3{font-size:15.5px;color:var(--navy)}
.flow-mod p{font-size:11.5px;color:var(--grey);line-height:1.4}
.flow-out{font-size:11px;color:var(--teal)}
.flow-arrow{flex:none;width:34px;text-align:center;color:#b8c9d2;font-size:19px}
.adapter-bar{margin-top:26px;margin-left:calc(25% + 10px);
  background:#e8f2f3;border:1px solid #a9d3d8;padding:20px 24px}
.adapter-name{font-size:15px;font-weight:700;color:var(--teal);display:block;margin-bottom:14px}
.pills{display:flex;gap:22px}
.pill{flex:1;background:#fff;border:1px solid #a9d3d8;padding:13px 10px;text-align:center;
  font-family:var(--mono);font-size:12.5px;font-weight:700;color:var(--navy)}
/* footnotes */
.footnote{font-size:13px;color:var(--grey);line-height:1.55;margin-top:20px;max-width:95%}
/* title + closing */
.slide--title,.slide--closing{background:var(--navy);justify-content:center;
  padding-left:110px;border-left:0}
.slide--title::before,.slide--closing::before{content:"";position:absolute;left:0;top:0;
  bottom:0;width:22px;background:var(--teal)}
.title-wrap h1{font-size:45px;color:#fff;line-height:1.14;letter-spacing:-.02em;margin:14px 0 0}
.subtitle{font-size:20px;color:#9fc5d6;margin-top:18px}
.rule{width:150px;height:3px;background:var(--teal);border:0;margin:30px 0}
.authors{font-size:17px;color:#fff;font-weight:700}
.affil{font-size:13px;color:#8fa8b8;margin-top:9px;line-height:1.55}
.h1--light{color:#fff;font-size:34px;margin-bottom:26px}
.thanks{font-size:21px;color:#fff;font-weight:700;margin-top:30px}
.ack{font-size:12px;color:#8fa8b8;margin-top:22px;line-height:1.6;max-width:86%}
/* chrome */
.slide-num{position:absolute;right:69px;bottom:22px;font-size:11px;color:#9aa8b0;
  font-weight:700}
.slide--title .slide-num,.slide--closing .slide-num{display:none}
#bar{position:fixed;left:0;top:0;height:3px;background:var(--teal);width:0;
  transition:width .18s ease;z-index:20}
#hint{position:fixed;right:16px;bottom:12px;color:#5d7285;font-size:11px;font-family:var(--font);
  z-index:20;transition:opacity .4s}
#hint.gone{opacity:0}
@media print{
  html,body{height:auto;overflow:visible;background:#fff}
  #deck{position:static;display:block}
  .slide{position:relative;display:flex!important;page-break-after:always;
    transform:none!important;margin:0 auto}
  #bar,#hint{display:none}
  @page{size:1280px 720px;margin:0}
}
</style>
</head>
<body>
<div id="bar"></div>
<div id="deck">
__BODY__
</div>
<div id="hint">← → navigate &nbsp;·&nbsp; F fullscreen &nbsp;·&nbsp; P print / PDF</div>
<script>
const slides=[...document.querySelectorAll('.slide')];
const notes=__NOTES__;
let i=0;
function fit(){
  const s=Math.min(innerWidth/1280,innerHeight/720);
  slides.forEach(el=>{el.style.transform='scale('+s+')'});
}
function show(n,fromHash){
  i=Math.max(0,Math.min(slides.length-1,n));
  slides.forEach((el,k)=>el.classList.toggle('is-active',k===i));
  document.querySelectorAll('.slide-num').forEach((el,k)=>el.textContent=(k+1));
  document.getElementById('bar').style.width=((i+1)/slides.length*100)+'%';
  if(!fromHash)history.replaceState(null,'','#'+(i+1));
}
addEventListener('hashchange',()=>{
  const n=parseInt(location.hash.slice(1),10);
  if(!isNaN(n))show(n-1,true);
});
addEventListener('resize',fit);
addEventListener('keydown',e=>{
  if(e.key==='ArrowRight'||e.key===' '||e.key==='PageDown'){show(i+1);e.preventDefault();}
  else if(e.key==='ArrowLeft'||e.key==='PageUp'){show(i-1);e.preventDefault();}
  else if(e.key==='Home')show(0); else if(e.key==='End')show(slides.length-1);
  else if(e.key==='f'||e.key==='F'){document.fullscreenElement?document.exitFullscreen():document.documentElement.requestFullscreen();}
  else if(e.key==='p'||e.key==='P')window.print();
  else if(e.key==='n'||e.key==='N')console.log('NOTES %d/%d\\n\\n%s',i+1,slides.length,notes[i]||'(none)');
});
let t=setTimeout(()=>document.getElementById('hint').classList.add('gone'),4200);
addEventListener('mousemove',()=>{const h=document.getElementById('hint');
  h.classList.remove('gone');clearTimeout(t);t=setTimeout(()=>h.classList.add('gone'),2600);});
fit();
const start=parseInt(location.hash.slice(1),10);
show(isNaN(start)?0:start-1,true);
</script>
</body>
</html>
"""


def build_html():
    import json
    notes = [s["notes"] for s in SLIDES]
    doc = HTML.replace("__BODY__", html_body()).replace(
        "__NOTES__", json.dumps(notes))
    with open(HTML_OUT, "w", encoding="utf-8") as f:
        f.write(doc)
    print(f"html  {HTML_OUT}  ({len(SLIDES)} slides, "
          f"{os.path.getsize(HTML_OUT)//1024}KB)")


if __name__ == "__main__":
    build_pptx()
    build_html()
